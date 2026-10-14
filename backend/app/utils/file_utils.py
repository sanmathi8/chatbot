import os
import re
import unicodedata
from pathlib import Path
from typing import List, Dict, Any, Tuple
from app.core.exceptions import InvalidFileTypeError, FileTooLargeError, DocumentProcessingError
from app.core.config import settings
from app.core.logging import logger

def sanitize_filename(filename: str) -> str:
    """Sanitize filename to prevent directory traversal and remove illegal chars."""
    clean = os.path.basename(filename)
    clean = unicodedata.normalize("NFKD", clean).encode("ascii", "ignore").decode("ascii")
    # Keep alphanumeric, dots, underscores, dashes
    clean = re.sub(r"[^a-zA-Z0-9._-]", "_", clean)
    clean = re.sub(r"_+", "_", clean)
    return clean if clean else "document.txt"

def validate_file(filename: str, file_size: int) -> str:
    """Validate file extension and file size."""
    ext = Path(filename).suffix.lower()
    if ext not in settings.ALLOWED_EXTENSIONS:
        raise InvalidFileTypeError(ext)
    
    max_bytes = settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024
    if file_size > max_bytes:
        raise FileTooLargeError(settings.MAX_UPLOAD_SIZE_MB)
    
    if file_size <= 0:
        raise DocumentProcessingError("The uploaded file is empty.")
    
    return ext

def clean_text(text: str) -> str:
    """Clean extra spaces while preserving paragraph breaks."""
    # Replace non-breaking spaces
    text = text.replace("\u00a0", " ")
    # Replace carriage returns
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Replace 3 or more consecutive newlines with 2
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Replace multiple spaces/tabs with single space
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()

def extract_text_from_pdf(file_path: Path) -> List[Dict[str, Any]]:
    """Extract text from PDF file with page attribution."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(file_path))
        pages_content = []
        
        for idx, page in enumerate(reader.pages):
            raw_page_text = page.extract_text() or ""
            cleaned = clean_text(raw_page_text)
            if cleaned:
                pages_content.append({
                    "text": cleaned,
                    "page_number": idx + 1
                })
        
        if not pages_content:
            raise DocumentProcessingError("No extractable text found in PDF (it may contain only scanned images).")
        
        return pages_content
    except Exception as e:
        if isinstance(e, DocumentProcessingError):
            raise
        logger.error(f"Failed to extract text from PDF: {e}")
        raise DocumentProcessingError(f"Error reading PDF file: {str(e)}")

def extract_text_from_docx(file_path: Path) -> List[Dict[str, Any]]:
    """Extract text from DOCX file."""
    try:
        from docx import Document
        doc = Document(str(file_path))
        full_text = []
        
        # Paragraphs
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text.strip())
                
        # Tables
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))
                    
        cleaned = clean_text("\n\n".join(full_text))
        if not cleaned:
            raise DocumentProcessingError("No extractable text found in DOCX file.")
            
        return [{"text": cleaned, "page_number": 1}]
    except Exception as e:
        if isinstance(e, DocumentProcessingError):
            raise
        logger.error(f"Failed to extract text from DOCX: {e}")
        raise DocumentProcessingError(f"Error reading DOCX file: {str(e)}")

def extract_text_from_txt(file_path: Path) -> List[Dict[str, Any]]:
    """Extract text from plain text or markdown file with multi-encoding fallback."""
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    content = None
    
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                content = f.read()
                break
        except UnicodeDecodeError:
            continue
            
    if content is None:
        raise DocumentProcessingError("Unable to decode text file with supported encodings.")
        
    cleaned = clean_text(content)
    if not cleaned:
        raise DocumentProcessingError("The text document contains no extractable text.")
        
    return [{"text": cleaned, "page_number": 1}]

def extract_document_text(file_path: Path, extension: str) -> List[Dict[str, Any]]:
    """Dispatch file extraction based on file extension."""
    ext = extension.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext in [".txt", ".md"]:
        return extract_text_from_txt(file_path)
    else:
        raise InvalidFileTypeError(ext)
